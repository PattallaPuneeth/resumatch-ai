import os
import tempfile
from collections.abc import Generator

import docx
import fitz
import pytest
from PIL import Image, ImageDraw

from ml.pipelines.intake.exceptions import (
    CorruptFileError,
    EncryptedPdfError,
    TextTooShortAfterOcrError,
    UnsupportedTypeError,
)
from ml.pipelines.intake.parser import parse_document


@pytest.fixture
def temp_dir() -> Generator[str, None, None]:
    with tempfile.TemporaryDirectory() as td:
        yield td


def create_sample_pdf(path: str, is_multi_column: bool = False) -> str:
    """Creates a real test PDF document with realistic resume content."""
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4

    if is_multi_column:
        # Left column
        left_text = (
            "Jane Doe\n"
            "jane.doe@example.com\n"
            "(555) 123-4567\n\n"
            "SKILLS\n"
            "Python, TypeScript, SQL\n"
            "FastAPI, React, Node.js\n"
            "Docker, Kubernetes, Redis\n"
            "PostgreSQL, Git, Linux\n"
        )
        page.insert_text((50, 72), left_text, fontsize=10)

        # Right column
        right_text = (
            "EXPERIENCE\n"
            "Senior Engineer at Acme Corp\n"
            "- Designed distributed queue pipelines using BullMQ.\n"
            "- Architected microservices with FastAPI and Docker.\n\n"
            "EDUCATION\n"
            "Bachelor of Science in Computer Science\n"
            "University of California, 2020\n"
        )
        page.insert_text((300, 72), right_text, fontsize=10)
    else:
        full_text = (
            "Jane Doe\n"
            "jane.doe@example.com\n"
            "(555) 123-4567\n"
            "https://linkedin.com/in/janedoe\n"
            "https://github.com/janedoe\n\n"
            "SUMMARY\n"
            "Experienced Senior Software Engineer specializing in Python, "
            "FastAPI, and cloud systems.\n\n"
            "SKILLS\n"
            "Languages: Python, TypeScript, JavaScript, SQL, HTML, CSS\n"
            "Frameworks: FastAPI, React, Node.js, Next.js, Django\n"
            "Tools: Docker, Kubernetes, PostgreSQL, Redis, MinIO, Git, Linux\n"
            "Soft Skills: Leadership, Communication, Problem Solving\n\n"
            "EXPERIENCE\n"
            "Senior Software Engineer at TechCorp\n"
            "- Architected scalable document parsing pipelines with PyMuPDF and BullMQ.\n"
            "- Improved system throughput by 45% using Redis caching.\n\n"
            "EDUCATION\n"
            "Bachelor of Science in Computer Science\n"
            "University of Technology, 2021\n"
        )
        page.insert_text((50, 72), full_text, fontsize=10)

    doc.save(path)
    doc.close()
    return path


def create_sample_docx(path: str, include_table: bool = True) -> str:
    """Creates a real test DOCX document with realistic resume content."""
    doc = docx.Document()
    doc.add_heading("Alex Rivera", level=1)
    doc.add_paragraph("alex.rivera@example.com | (555) 987-6543")
    doc.add_paragraph("https://linkedin.com/in/alexrivera")

    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Staff Backend Engineer with over 8 years of experience building high-concurrency systems."
    )

    doc.add_heading("Skills", level=2)
    doc.add_paragraph("Python, Go, PostgreSQL, Docker, Kubernetes, Redis, AWS, Git")

    doc.add_heading("Experience", level=2)
    doc.add_paragraph("Lead Backend Architect - CloudScale")
    doc.add_paragraph("- Built asynchronous event pipelines processing millions of daily events.")
    doc.add_paragraph("- Spearheaded cloud migration to Kubernetes on AWS.")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("Master of Science in Software Engineering, State University, 2018")

    if include_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Certification"
        table.cell(0, 1).text = "Year"
        table.cell(1, 0).text = "AWS Certified Solutions Architect"
        table.cell(1, 1).text = "2022"

    doc.save(path)
    return path


def create_scanned_image_pdf(path: str, text: str) -> str:
    """Creates an image-only PDF containing no selectable text layer."""
    # Render text onto an image using Pillow
    img = Image.new("RGB", (600, 800), color=(255, 255, 255))
    draw = ImageDraw.Draw(img)
    draw.text((50, 50), text, fill=(0, 0, 0))

    # Save image as PDF using PyMuPDF
    doc = fitz.open()
    page = doc.new_page(width=600, height=800)
    img_byte_arr = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
    try:
        img.save(img_byte_arr.name)
        page.insert_image(page.rect, filename=img_byte_arr.name)
        doc.save(path)
    finally:
        doc.close()
        if os.path.exists(img_byte_arr.name):
            os.remove(img_byte_arr.name)

    return path


def test_valid_pdf_extraction(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "resume.pdf")
    create_sample_pdf(pdf_path, is_multi_column=False)

    result = parse_document(pdf_path, "application/pdf")
    assert result.extraction_method == "TEXT_LAYER"
    assert result.page_count == 1
    assert "Jane Doe" in result.raw_text
    assert result.contact.email == "jane.doe@example.com"
    assert result.contact.phone == "(555) 123-4567"
    assert len(result.skills) > 0
    assert any(s.name == "Python" for s in result.skills)
    assert any(s.name == "FastAPI" for s in result.skills)
    assert result.ats_score >= 80.0

    # Verify exact character offsets into canonical raw_text
    for skill in result.skills:
        assert skill.confidence is not None
        span = skill.confidence.source_span
        assert span is not None
        assert result.raw_text[span.start_char : span.end_char].lower() == skill.name.lower()

    if result.contact.email_confidence:
        span = result.contact.email_confidence.source_span
        assert span is not None
        assert result.raw_text[span.start_char : span.end_char] == result.contact.email


def test_valid_docx_extraction(temp_dir: str) -> None:
    docx_path = os.path.join(temp_dir, "resume.docx")
    create_sample_docx(docx_path, include_table=True)

    result = parse_document(
        docx_path,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert result.extraction_method == "DOCX"
    assert "Alex Rivera" in result.raw_text
    assert result.contact.email == "alex.rivera@example.com"
    assert any(s.name == "PostgreSQL" for s in result.skills)
    # Tables should be flagged as an ATS hazard
    assert any(h.code == "TABLES_PRESENT" for h in result.ats_hazards)


def test_multi_column_pdf_detection(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "multi_column.pdf")
    create_sample_pdf(pdf_path, is_multi_column=True)

    result = parse_document(pdf_path, "application/pdf")
    assert result.extraction_method == "TEXT_LAYER"
    assert any(h.code == "MULTI_COLUMN_LAYOUT" for h in result.ats_hazards)
    assert result.ats_score < 100.0


def test_scanned_pdf_triggers_ocr(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "scanned.pdf")
    scanned_text = (
        "Thomas Edison\n"
        "thomas.edison@example.com\n"
        "(555) 444-3322\n\n"
        "EXPERIENCE\n"
        "Research Scientist at Edison Labs\n"
        "- Developed phonograph and electric power distribution systems.\n"
        "- Authored over one thousand patented innovations.\n\n"
        "SKILLS\n"
        "Physics, Electrical Engineering, Experimentation, Leadership\n\n"
        "EDUCATION\n"
        "Self-taught Academic, 1870\n"
    )
    create_scanned_image_pdf(pdf_path, scanned_text)

    status_transitions: list[str] = []

    def status_cb(status: str) -> None:
        status_transitions.append(status)

    result = parse_document(pdf_path, "application/pdf", status_callback=status_cb)
    assert "NEEDS_OCR" in status_transitions
    assert "OCR_PROCESSING" in status_transitions
    assert result.extraction_method == "OCR"
    assert any(h.code == "OCR_REQUIRED" for h in result.ats_hazards)
    assert "Edison" in result.raw_text


def test_ocr_text_too_short_fails(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "too_short.pdf")
    # Tiny text (< 200 chars)
    create_scanned_image_pdf(pdf_path, "Short scan")

    with pytest.raises(TextTooShortAfterOcrError):
        parse_document(pdf_path, "application/pdf")


def test_encrypted_pdf_fails(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "encrypted.pdf")
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Confidential Resume Data")
    doc.save(pdf_path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="secret123")
    doc.close()

    with pytest.raises(EncryptedPdfError):
        parse_document(pdf_path, "application/pdf")


def test_corrupt_pdf_fails(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "corrupt.pdf")
    with open(pdf_path, "wb") as f:
        f.write(b"%PDF-1.4 garbage corrupt bytes not a valid pdf header \x00\xff")

    with pytest.raises(CorruptFileError):
        parse_document(pdf_path, "application/pdf")


def test_unsupported_type_fails(temp_dir: str) -> None:
    txt_path = os.path.join(temp_dir, "resume.txt")
    with open(txt_path, "w") as f:
        f.write("Plain text resume")

    with pytest.raises(UnsupportedTypeError):
        parse_document(txt_path, "text/plain")


def test_empty_zero_byte_file_fails(temp_dir: str) -> None:
    empty_path = os.path.join(temp_dir, "empty.pdf")
    with open(empty_path, "wb") as f:
        f.write(b"")

    with pytest.raises(CorruptFileError):
        parse_document(empty_path, "application/pdf")


def test_textless_blank_pdf_fails_ocr_too_short(temp_dir: str) -> None:
    pdf_path = os.path.join(temp_dir, "blank.pdf")
    doc = fitz.open()
    doc.new_page()  # Completely blank white page
    doc.save(pdf_path)
    doc.close()

    with pytest.raises(TextTooShortAfterOcrError):
        parse_document(pdf_path, "application/pdf")


def test_corrupt_docx_fails(temp_dir: str) -> None:
    docx_path = os.path.join(temp_dir, "corrupt.docx")
    with open(docx_path, "wb") as f:
        f.write(b"PK\x03\x04 corrupt zip payload not a real word document")

    with pytest.raises(CorruptFileError):
        parse_document(
            docx_path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


def test_docx_zip_bomb_fails_safely(temp_dir: str) -> None:
    """Simulates a zip bomb with extreme compression ratio."""
    import zipfile
    bomb_path = os.path.join(temp_dir, "zipbomb.docx")
    # Highly compressible stream of 1MB zeroes compressed into tiny bytes
    with zipfile.ZipFile(bomb_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * 1000000)

    with pytest.raises(CorruptFileError, match="potential zip bomb"):
        parse_document(
            bomb_path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


def test_docx_with_macro_payload_fails(temp_dir: str) -> None:
    """Verifies that DOCX with embedded VBA macro is rejected."""
    import zipfile
    macro_path = os.path.join(temp_dir, "macro.docx")
    with zipfile.ZipFile(macro_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"<xml>Dummy text</xml>")
        zf.writestr("word/vbaProject.bin", b"fake macro binary payload")

    with pytest.raises(CorruptFileError, match="macros or executable payloads"):
        parse_document(
            macro_path,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


def test_pdf_with_oversized_dimensions_fails(temp_dir: str) -> None:
    """Verifies that PDF with massive dimensions (e.g. 5000x5000 pt) is rejected."""
    pdf_path = os.path.join(temp_dir, "huge_dim.pdf")
    doc = fitz.open()
    doc.new_page(width=5000, height=5000)
    doc.save(pdf_path)
    doc.close()

    with pytest.raises(CorruptFileError, match="exceed maximum safety limit"):
        parse_document(pdf_path, "application/pdf")

